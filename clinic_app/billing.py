from __future__ import annotations

import datetime as dt
import secrets
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.templating import Jinja2Templates
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import or_
from sqlalchemy.orm import Session
from docx import Document

from clinic_app.db import get_db
from clinic_app.deps import CurrentUser, forbid_demo_writes, get_current_user
from clinic_app.models import Bill, BillLine, Patient, Product, StockBatch
from clinic_app.settings import settings


router = APIRouter(prefix="/billing")
templates = Jinja2Templates(directory="clinic_app/templates")


def _d(x: str | None, default: Decimal = Decimal("0")) -> Decimal:
    if x is None:
        return default
    s = str(x).strip()
    if s == "":
        return default
    try:
        return Decimal(s)
    except InvalidOperation:
        return default


def _bill_or_404(db: Session, *, tenant_id: int, bill_id: int) -> Bill:
    bill = (
        db.query(Bill)
        .filter(Bill.tenant_id == tenant_id)
        .filter(Bill.id == bill_id)
        .first()
    )
    if bill is None:
        raise HTTPException(status_code=404, detail="Bill not found")
    return bill


def _deduct_stock_fifo(db: Session, *, tenant_id: int, product_id: int, qty_needed: int) -> None:
    if qty_needed <= 0:
        return
    batches = (
        db.query(StockBatch)
        .filter(StockBatch.tenant_id == tenant_id)
        .filter(StockBatch.product_id == product_id)
        .filter(StockBatch.qty_on_hand > 0)
        .order_by(StockBatch.expiry_date.asc().nullslast(), StockBatch.id.asc())
        .all()
    )
    remaining = qty_needed
    for b in batches:
        if remaining <= 0:
            break
        take = min(int(b.qty_on_hand), remaining)
        b.qty_on_hand = int(b.qty_on_hand) - take
        remaining -= take
    # If remaining > 0, we allow negative stock? No—keep it at 0; later we’ll surface warnings in UI.


@router.get("", response_class=HTMLResponse)
def bills_list(
    request: Request,
    q: str | None = None,
    user: CurrentUser = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(Bill).filter(Bill.tenant_id == user.tenant_id)
    if q and q.strip():
        s = f"%{q.strip()}%"
        query = query.join(Patient, Patient.id == Bill.patient_id).filter(
            or_(Bill.bill_no.ilike(s), Patient.name.ilike(s))
        )
    bills = query.order_by(Bill.bill_date.desc(), Bill.id.desc()).limit(200).all()
    return templates.TemplateResponse(
        "bills_list.html",
        {"request": request, "user": user, "bills": bills, "q": q or ""},
    )


@router.get("/new", response_class=HTMLResponse)
def new_bill_page(
    request: Request,
    user: CurrentUser = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    patients = (
        db.query(Patient)
        .filter(Patient.tenant_id == user.tenant_id)
        .order_by(Patient.name.asc())
        .all()
    )
    products = (
        db.query(Product)
        .filter(Product.tenant_id == user.tenant_id)
        .order_by(Product.name.asc())
        .all()
    )
    return templates.TemplateResponse(
        "bill_form.html",
        {
            "request": request,
            "user": user,
            "patients": patients,
            "products": products,
            "today": dt.date.today(),
            "clinic_name": settings.clinic_name_default,
            "clinic_phone": settings.clinic_phone_default,
            "clinic_address": settings.clinic_address_default,
        },
    )


@router.post("/new")
def create_bill(
    patient_id: int = Form(...),
    bill_date: str | None = Form(default=None),
    clinic_name: str = Form(...),
    clinic_phone: str | None = Form(default=None),
    clinic_address: str | None = Form(default=None),
    consulting_fee: str | None = Form(default=None),
    line_type: list[str] = Form(default=[]),
    description: list[str] = Form(default=[]),
    qty: list[str] = Form(default=[]),
    unit_price: list[str] = Form(default=[]),
    gst_rate: list[str] = Form(default=[]),
    product_id: list[str] = Form(default=[]),
    user: CurrentUser = Depends(forbid_demo_writes),
    db: Session = Depends(get_db),
):
    bdate = dt.date.today()
    if bill_date:
        try:
            bdate = dt.date.fromisoformat(bill_date)
        except ValueError:
            bdate = dt.date.today()

    patient = (
        db.query(Patient)
        .filter(Patient.tenant_id == user.tenant_id)
        .filter(Patient.id == patient_id)
        .first()
    )
    if patient is None:
        raise HTTPException(status_code=400, detail="Invalid patient")

    tmp_no = f"TMP-{secrets.token_hex(4)}"
    consult = _d(consulting_fee)

    bill = Bill(
        tenant_id=user.tenant_id,
        patient_id=patient.id,
        bill_no=tmp_no,
        bill_date=bdate,
        consulting_fee=float(consult),
        subtotal=0,
        gst_total=0,
        total=0,
        clinic_name=clinic_name.strip(),
        clinic_phone=(clinic_phone.strip() if clinic_phone else ""),
        clinic_address=(clinic_address.strip() if clinic_address else ""),
    )
    db.add(bill)
    db.flush()
    bill.bill_no = f"B{bill.id:05d}"

    subtotal = Decimal("0") + consult
    gst_total_d = Decimal("0")
    total_d = Decimal("0") + consult

    n = min(len(line_type), len(description), len(qty), len(unit_price), len(gst_rate), len(product_id))
    for i in range(n):
        lt = (line_type[i] or "").strip()
        desc = (description[i] or "").strip()
        if not lt or not desc:
            continue

        q = _d(qty[i], default=Decimal("1"))
        up = _d(unit_price[i])
        gstp = _d(gst_rate[i])
        line_sub = (q * up).quantize(Decimal("0.01"))
        line_gst = (line_sub * gstp / Decimal("100")).quantize(Decimal("0.01"))
        line_total = (line_sub + line_gst).quantize(Decimal("0.01"))

        pid = None
        if lt == "medicine":
            try:
                pid_int = int(str(product_id[i]).strip() or "0")
                pid = pid_int if pid_int > 0 else None
            except ValueError:
                pid = None

        line = BillLine(
            bill_id=bill.id,
            tenant_id=user.tenant_id,
            product_id=pid,
            line_type=lt,
            description=desc[:240],
            qty=float(q),
            unit_price=float(up),
            gst_rate=float(gstp),
            gst_amount=float(line_gst),
            line_total=float(line_total),
        )
        db.add(line)

        subtotal += line_sub
        gst_total_d += line_gst
        total_d += line_total

        if lt == "medicine" and pid is not None:
            # Deduct stock using FIFO by expiry.
            _deduct_stock_fifo(db, tenant_id=user.tenant_id, product_id=pid, qty_needed=int(q))

    bill.subtotal = float(subtotal)
    bill.gst_total = float(gst_total_d)
    bill.total = float(total_d)
    db.commit()

    return RedirectResponse(url=f"/billing/{bill.id}", status_code=302)


@router.get("/{bill_id}", response_class=HTMLResponse)
def bill_detail(
    request: Request,
    bill_id: int,
    user: CurrentUser = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    bill = _bill_or_404(db, tenant_id=user.tenant_id, bill_id=bill_id)
    lines = (
        db.query(BillLine)
        .filter(BillLine.tenant_id == user.tenant_id)
        .filter(BillLine.bill_id == bill.id)
        .order_by(BillLine.id.asc())
        .all()
    )
    return templates.TemplateResponse(
        "bill_detail.html",
        {"request": request, "user": user, "bill": bill, "lines": lines},
    )


@router.get("/{bill_id}/export/pdf")
def export_bill_pdf(
    bill_id: int,
    user: CurrentUser = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    bill = _bill_or_404(db, tenant_id=user.tenant_id, bill_id=bill_id)
    lines = (
        db.query(BillLine)
        .filter(BillLine.tenant_id == user.tenant_id)
        .filter(BillLine.bill_id == bill.id)
        .order_by(BillLine.id.asc())
        .all()
    )

    buf = bytearray()
    from io import BytesIO

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, title=f"Bill {bill.bill_no}")
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(f"<b>{bill.clinic_name}</b>", styles["Title"]))
    if bill.clinic_address:
        story.append(Paragraph(bill.clinic_address.replace("\n", "<br/>"), styles["Normal"]))
    if bill.clinic_phone:
        story.append(Paragraph(f"Phone: {bill.clinic_phone}", styles["Normal"]))
    story.append(Spacer(1, 10))
    story.append(Paragraph(f"<b>Bill No:</b> {bill.bill_no} &nbsp;&nbsp; <b>Date:</b> {bill.bill_date.isoformat()}", styles["Normal"]))
    story.append(Paragraph(f"<b>Patient:</b> {bill.patient.name}", styles["Normal"]))
    story.append(Spacer(1, 12))

    data = [["Type", "Description", "Qty", "Unit", "GST%", "Amount"]]
    if float(bill.consulting_fee) > 0:
        data.append(["consulting", "Consulting fee", "1", "-", "0", f"{float(bill.consulting_fee):.2f}"])
    for ln in lines:
        data.append(
            [
                ln.line_type,
                ln.description,
                f"{float(ln.qty):g}",
                f"{float(ln.unit_price):.2f}",
                f"{float(ln.gst_rate):g}",
                f"{float(ln.line_total):.2f}",
            ]
        )

    tbl = Table(data, colWidths=[70, 250, 45, 55, 45, 70])
    tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 1), (-1, -1), colors.whitesmoke),
            ]
        )
    )
    story.append(tbl)
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"<b>Subtotal:</b> ₹{float(bill.subtotal):.2f}", styles["Normal"]))
    story.append(Paragraph(f"<b>GST Total:</b> ₹{float(bill.gst_total):.2f}", styles["Normal"]))
    story.append(Paragraph(f"<b>Total:</b> ₹{float(bill.total):.2f}", styles["Heading2"]))

    doc.build(story)
    pdf_bytes = out.getvalue()

    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{bill.bill_no}.pdf"'},
    )


@router.get("/{bill_id}/export/docx")
def export_bill_docx(
    bill_id: int,
    user: CurrentUser = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    bill = _bill_or_404(db, tenant_id=user.tenant_id, bill_id=bill_id)
    lines = (
        db.query(BillLine)
        .filter(BillLine.tenant_id == user.tenant_id)
        .filter(BillLine.bill_id == bill.id)
        .order_by(BillLine.id.asc())
        .all()
    )

    doc = Document()
    doc.add_heading(bill.clinic_name, level=1)
    if bill.clinic_address:
        doc.add_paragraph(bill.clinic_address)
    if bill.clinic_phone:
        doc.add_paragraph(f"Phone: {bill.clinic_phone}")

    doc.add_paragraph(f"Bill No: {bill.bill_no}    Date: {bill.bill_date.isoformat()}")
    doc.add_paragraph(f"Patient: {bill.patient.name}")

    table = doc.add_table(rows=1, cols=6)
    hdr = table.rows[0].cells
    hdr[0].text = "Type"
    hdr[1].text = "Description"
    hdr[2].text = "Qty"
    hdr[3].text = "Unit"
    hdr[4].text = "GST%"
    hdr[5].text = "Amount"

    if float(bill.consulting_fee) > 0:
        r = table.add_row().cells
        r[0].text = "consulting"
        r[1].text = "Consulting fee"
        r[2].text = "1"
        r[3].text = "-"
        r[4].text = "0"
        r[5].text = f"{float(bill.consulting_fee):.2f}"

    for ln in lines:
        r = table.add_row().cells
        r[0].text = ln.line_type
        r[1].text = ln.description
        r[2].text = f"{float(ln.qty):g}"
        r[3].text = f"{float(ln.unit_price):.2f}"
        r[4].text = f"{float(ln.gst_rate):g}"
        r[5].text = f"{float(ln.line_total):.2f}"

    doc.add_paragraph(f"Subtotal: ₹{float(bill.subtotal):.2f}")
    doc.add_paragraph(f"GST Total: ₹{float(bill.gst_total):.2f}")
    doc.add_paragraph(f"Total: ₹{float(bill.total):.2f}")

    from io import BytesIO

    out = BytesIO()
    doc.save(out)
    data = out.getvalue()
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{bill.bill_no}.docx"'},
    )

